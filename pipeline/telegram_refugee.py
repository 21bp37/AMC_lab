"""Filtr uchodźczy dla Telegrama (RU/UK): diagnostyka gęstości szeregu i ręczne sprawdzenie trafności.

Obecny filtr (Ukraina ORAZ Polska) łapie głównie wojnę i geopolitykę. Filtr uchodźczy dokłada trzeci warunek: post mówi
o uchodźcach / migrantach / Ukraińcach w Polsce (słowa w pipeline/config.py: TELEGRAM_REFUGEE_TERMS). Działa na już
pobranych i przefiltrowanych plikach (posty UA+PL), więc NIE trzeba pobierać 49 GB od nowa. Opis: reports/FILTR_UCHODZCZY.md.

Dwa podpolecenia:

  check  - liczy, ile postów zostaje i czy dzienny szereg jest dość gęsty, zapisuje szereg dzienny i losuje próbę do przeglądu
           (posty PRZEPUSZCZONE i ODRZUCONE przez filtr) jako arkusz XLSX/CSV z tłumaczeniem na polski.
           python -m pipeline.telegram_refugee check --translator google
  score  - po wypełnieniu kolumny „o_uchodzcach” (tak / nie / niejasne) liczy precyzję i szacowaną czułość filtra.
           python -m pipeline.telegram_refugee score

Co sprawdza się automatycznie (check): gęstość, udział zatrzymanych postów, największe dni, najczęstsze trafione rdzenie.
Czego automat nie sprawdzi (score, trzeba ręcznie ok. 80 postów): czy zatrzymane posty naprawdę są o uchodźcach (precyzja)
i ilu uchodźczych postów filtr gubi (czułość).
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd

from pipeline import annotation_sample as sm
from pipeline import config, eda_telegram

SIDES = ["Pro-Russian", "Pro-Ukrainian"]
CHOICES = ["tak", "nie", "niejasne"]
COMPILED = {k: re.compile(v) for k, v in config.TELEGRAM_REFUGEE.items()}
ANNOT = sm.ANNOT_DIR
REPORTS = config.REPORTS

REVIEW_RULES = [
    "Pytanie do każdego wiersza: czy post (po tłumaczeniu) dotyczy uchodźców z Ukrainy, migrantów albo Ukraińców mieszkających w Polsce "
    "(ich sytuacji, pobytu, traktowania, pomocy dla nich lub wrogości wobec nich)?",
    "tak = ludzie (uchodźcy / migranci / Ukraińcy w Polsce) są tematem posta, także gdy post ich atakuje albo broni.",
    "nie = wojna, broń, granice państw, polityka międzynarodowa, gospodarka itp. bez wątku ludzi; wzmianka o uchodźcach tylko w przelocie.",
    "niejasne = nie da się ocenić (zły przekład, urwany tekst). Opisz powód w kolumnie komentarz.",
    "Nie oceniasz, czy post jest „pro” czy „anty”. Nie zgaduj, czy filtr post przepuścił; kolumna z jego decyzją jest ukryta celowo.",
]


# --------------------------------------------------------------------------- filtr
def flag(msg: pd.Series) -> pd.DataFrame:
    """Flagi kategorii filtra uchodźczego i trafione rdzenie (do diagnostyki fałszywych trafień).

    refugee_weak = jakiekolwiek trafienie (wersja 0.1); is_refugee = „mocne” trafienie (od 0.2):
    fraza „Ukraińcy w Polsce” ALBO >= TELEGRAM_REFUGEE_MIN_HITS trafień ALBO pierwsze trafienie w pierwszych TELEGRAM_REFUGEE_LEAD_CHARS znaków."""
    m = msg.fillna("").astype(str)
    out = pd.DataFrame({k: m.str.contains(v, regex=True) for k, v in config.TELEGRAM_REFUGEE.items()}).astype(bool)
    weak = out.any(axis=1)
    strong = pd.Series(False, index=m.index)
    hits = pd.Series("", index=m.index, dtype=object)
    for i in out.index[weak]:
        spans = []
        for k, rx in COMPILED.items():
            for mt in rx.finditer(m.at[i]):
                spans.append((mt.start(), mt.group(0).lower()[:40]))
        spans.sort()
        strong.at[i] = bool(out.at[i, "ukraincy_w_polsce"] or len(spans) >= config.TELEGRAM_REFUGEE_MIN_HITS
                            or spans[0][0] < config.TELEGRAM_REFUGEE_LEAD_CHARS)
        hits.at[i] = "; ".join(list(dict.fromkeys(t for _, t in spans))[:6])
    out["refugee_weak"] = weak
    out["is_refugee"] = strong
    out["refugee_hits"] = hits
    return out


def load(posts_glob: str, clusters_csv: str) -> pd.DataFrame:
    sm.log("Wczytuję pliki z postami ...")
    raw = sm.load_posts_verbose(posts_glob)
    sm.log(f"Razem {sm.num(len(raw))} postów; okno, filtr Ukraina+Polska, strony kanałów ...")
    df = eda_telegram.prepare(raw, clusters_csv)               # okno, Ukraina+Polska, strona kanału, is_anti, copy_key
    sm.log(f"Po filtrze UA+PL: {sm.num(len(df))} postów; stosuję filtr uchodźczy ...")
    fl = flag(df["message"])
    return pd.concat([df.reset_index(drop=True), fl.reset_index(drop=True)], axis=1)


# --------------------------------------------------------------------------- gęstość
def _series_stats(sub: pd.DataFrame, idx: pd.DatetimeIndex) -> dict:
    s = sub.groupby("day").size().reindex(idx, fill_value=0)
    w = s.resample("W-SUN").sum()
    pk = s.idxmax()
    ev = [v for k, v in config.EVENTS.items() if abs((pd.Timestamp(k) - pk).days) <= 1]
    return {"median_day": float(s.median()), "mean_day": float(s.mean()), "zero_days": float((s == 0).mean()),
            "ge3_days": float((s >= 3).mean()), "mean_week": float(w.mean()), "median_week": float(w.median()),
            "peak_day": str(pk.date()), "peak_n": int(s.max()), "peak_event": "; ".join(ev)}


def density_verdict(st: dict, med_min: float = 5, zero_max: float = 0.30, week_min: float = 20) -> str:
    if st["median_day"] >= med_min and st["zero_days"] <= zero_max:
        return "dzienny szereg wystarczająco gęsty"
    if st["mean_week"] >= week_min:
        return "za rzadki na test dzienny; wystarczy test tygodniowy"
    return "za mało postów: szeregu nie da się sensownie zbudować (zawęzić projekt)"


def side_stats(d: pd.DataFrame, idx: pd.DatetimeIndex) -> dict:
    ref = d[d["is_refugee"]]
    top10 = ref["peer_id"].value_counts().head(10).sum() / len(ref) if len(ref) else float("nan")
    out = {"n_uapl": int(len(d)), "n_ref": int(len(ref)), "retained": len(ref) / len(d) if len(d) else float("nan"),
           "channels_ref": int(ref["peer_id"].nunique()), "top10_channel_share": float(top10),
           "copy_share": float(ref.duplicated("copy_key").mean()) if len(ref) else float("nan"),
           "anti_share_uapl": float(d["is_anti"].mean()) if len(d) else float("nan"),
           "anti_share_ref": float(ref["is_anti"].mean()) if len(ref) else float("nan")}
    out["ref"] = _series_stats(ref, idx)
    out["ref_anti"] = _series_stats(ref[ref["is_anti"]], idx)
    out["ref"]["verdict"] = density_verdict(out["ref"])
    out["ref_anti"]["verdict"] = density_verdict(out["ref_anti"])
    return out


def daily_table(df: pd.DataFrame, idx: pd.DatetimeIndex) -> pd.DataFrame:
    rows = []
    for side in SIDES:
        d = df[df["side"] == side]
        ref = d[d["is_refugee"]]
        t = pd.DataFrame({"n_topic": d.groupby("day").size(), "n_refugee": ref.groupby("day").size(),
                          "n_refugee_anti": ref[ref["is_anti"]].groupby("day").size()}).reindex(idx, fill_value=0).fillna(0).astype(int)
        t.insert(0, "side", side)
        t.index.name = "day"
        rows.append(t.reset_index())
    return pd.concat(rows, ignore_index=True)


def pct(x, nd=0):
    return "–" if x is None or (isinstance(x, float) and math.isnan(x)) else f"{100 * x:.{nd}f}%"


def results_md(stats: dict, cats: dict, top_hits: list, version: str) -> str:
    L = [f"# Filtr uchodźczy RU/UK, wyniki diagnostyki (wersja słów {version})", "",
         "Wygenerowane przez `python -m pipeline.telegram_refugee check`. Dotyczy postów UA+PL z pobranych plików; "
         "okno 4.03–31.12.2022 (303 dni). Opis działania: `FILTR_UCHODZCZY.md`.", ""]
    for side in SIDES:
        s = stats.get(side)
        if not s or not s["n_uapl"]:
            continue
        r, ra = s["ref"], s["ref_anti"]
        L += [f"## {side}", "",
              f"- Posty UA+PL: {s['n_uapl']}; po filtrze uchodźczym: {s['n_ref']} ({pct(s['retained'])}), z {s['channels_ref']} kanałów; "
              f"10 największych kanałów daje {pct(s['top10_channel_share'])}; kopie tego samego tekstu: {pct(s['copy_share'])}.",
              f"- Udział postów z narracją anty-UA (słownik RU/UK v0.1): przed filtrem {pct(s['anti_share_uapl'], 1)}, po filtrze {pct(s['anti_share_ref'], 1)}. "
              "Filtr nie używa słów ze słownika anty-UA, więc różnica nie wynika z tych samych słów, ale może odzwierciedlać to, że posty o uchodźcach "
              "częściej dotyczą świadczeń itp. albo styl kanałów (zob. FILTR_UCHODZCZY.md).", "",
              "| szereg | mediana/dzień | dni z zerem | dni z ≥3 | średnio/tydzień | najwyższy dzień | werdykt |", "|---|---|---|---|---|---|---|",
              f"| posty uchodźcze | {r['median_day']:.0f} | {pct(r['zero_days'])} | {pct(r['ge3_days'])} | {r['mean_week']:.0f} | "
              f"{r['peak_day']} ({r['peak_n']}){' – ' + r['peak_event'] if r['peak_event'] else ''} | {r['verdict']} |",
              f"| uchodźcze i anty-UA | {ra['median_day']:.0f} | {pct(ra['zero_days'])} | {pct(ra['ge3_days'])} | {ra['mean_week']:.0f} | "
              f"{ra['peak_day']} ({ra['peak_n']}){' – ' + ra['peak_event'] if ra['peak_event'] else ''} | {ra['verdict']} |", ""]
    L += ["## Które kategorie i rdzenie łapią posty (wszystkie strony)", "", "| kategoria | postów |", "|---|---|"]
    L += [f"| {k} | {v} |" for k, v in cats.items()]
    L += ["", "Najczęściej trafione fragmenty (do przeglądu: jeśli któryś dominuje i wygląda na szum, usuń go z listy):", "",
          "| fragment | postów |", "|---|---|"] + [f"| {w} | {c} |" for w, c in top_hits] + [""]
    L += ["Progi werdyktu są orientacyjne (mediana ≥5 postów dziennie i ≤30% dni z zerem dla testu dziennego; ≥20 tygodniowo dla tygodniowego). "
          "Gęstość to nie to samo co trafność: czy posty są o uchodźcach, sprawdza dopiero ręczny przegląd (`score`).", ""]
    return "\n".join(L)


# --------------------------------------------------------------------------- próba do przeglądu
def pools(df: pd.DataFrame) -> dict:
    """Pule unikalnych tekstów: (strona, grupa) -> ramka. Grupa: przepuszczone / odrzucone przez filtr."""
    d = df[df["side"].isin(SIDES)].copy()
    d = d[d["message"].fillna("").map(sm.content_len) >= sm.MIN_CHARS_TG]
    d = d.sort_values("ts").drop_duplicates("copy_key")
    return {(side, grp): d[(d["side"] == side) & (d["is_refugee"] == (grp == "przepuszczony"))]
            for side in SIDES for grp in ("przepuszczony", "odrzucony")}


def reflag_fn(posts_glob: str, clusters_csv: str):
    """Funkcja master -> nowe flagi filtra dla wierszy Telegrama (klucz: peer_id/id), liczone na pełnych tekstach z danych."""
    def fn(master: pd.DataFrame) -> pd.DataFrame:
        tg = master[master["platform"] == "Telegram"]
        df = load(posts_glob, clusters_csv)
        key = df["peer_id"].astype(str) + "/" + df["id"].astype(str)
        by_key = pd.DataFrame({"refugee_filter": np.where(df["is_refugee"], "tak", "nie"), "refugee_hits": df["refugee_hits"].to_numpy()},
                              index=key.to_numpy())
        by_key = by_key[~by_key.index.duplicated()]
        found = tg["source_id"].isin(by_key.index)
        if not found.all():
            sm.log(f"UWAGA: {int((~found).sum())} postów z próby nie znaleziono w danych; zostaje ich stary wynik filtra")
        tg = tg[found]
        out = by_key.loc[tg["source_id"].to_numpy()].copy()
        out.index = tg.index
        return out
    return fn


def draw_review(df: pd.DataFrame, n_pass: int, n_rej: int, seed: int) -> pd.DataFrame:
    parts = []
    for (side, grp), pool in pools(df).items():
        n = (n_pass if grp == "przepuszczony" else n_rej) // len(SIDES)
        n = min(n, len(pool))
        if n == 0:
            continue
        s = pool.sample(n=n, random_state=seed + len(parts)).copy()
        s["group"], s["pool_N"], s["n"] = grp, len(pool), n
        s["weight"] = len(pool) / n
        parts.append(s)
    if not parts:
        raise SystemExit("Brak postów do próby przeglądu (puste dane Telegrama?).")
    m = pd.concat(parts, ignore_index=True)
    m = m.sample(frac=1.0, random_state=seed + 99).reset_index(drop=True)
    m["lex_hits_raw"] = m["refugee_hits"]
    m["text_original"] = m["message"].map(lambda t: sm.clip(sm.anonymize(t)))
    m.insert(0, "item_id", [f"R{i:03d}" for i in range(1, len(m) + 1)])
    m["text_pl"] = ""
    m["translation_status"] = "brak"
    keep = ["item_id", "group", "side", "day", "peer_id", "refugee_hits", "is_anti", "pool_N", "n", "weight",
            "text_original", "text_pl", "translation_status"]
    return m[keep]


def review_view(master: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame({
        "id": master["item_id"],
        "tekst_pl": np.where(master["text_pl"] == "", "(brak tłumaczenia, patrz tekst_oryg)", master["text_pl"]),
        "tekst_oryg": master["text_original"], "o_uchodzcach": "", "tlumaczenie": "", "komentarz": ""})


def write_review_xlsx(path: Path, view: pd.DataFrame) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Przeglad"
    ws.append(list(view.columns))
    for row in view.itertuples(index=False):
        ws.append(list(row))
    widths = {"id": 8, "tekst_pl": 75, "tekst_oryg": 55, "o_uchodzcach": 14, "tlumaczenie": 13, "komentarz": 35}
    for i, col in enumerate(view.columns, start=1):
        c = ws.cell(row=1, column=i)
        ws.column_dimensions[c.column_letter].width = widths.get(col, 15)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    wrap = Alignment(wrap_text=True, vertical="top")
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(view.columns) + 1):
            ws.cell(row=r, column=c).alignment = wrap
        ws.cell(row=r, column=4).fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "C2"
    dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True, showErrorMessage=True)
    dv2 = DataValidation(type="list", formula1='"ok,słabe"', allow_blank=True)
    ws.add_data_validation(dv); ws.add_data_validation(dv2)
    dv.add(f"D2:D{ws.max_row}"); dv2.add(f"E2:E{ws.max_row}")
    ins = wb.create_sheet("Instrukcja")
    ins.column_dimensions["A"].width = 4
    ins.column_dimensions["B"].width = 130
    for i, r in enumerate(REVIEW_RULES, 1):
        ins.append([i, r])
    for row in ins.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


# --------------------------------------------------------------------------- ocena przeglądu
def wilson(k: int, n: int, z: float = 1.96) -> tuple[float, float, float]:
    if n == 0:
        return (float("nan"),) * 3
    p = k / n
    den = 1 + z * z / n
    c = (p + z * z / (2 * n)) / den
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return p, max(0.0, c - h), min(1.0, c + h)


def recall_ci(n_pass_pool: float, k_pass: int, n_pass: int, n_rej_pool: float, k_rej: int, n_rej: int,
              draws: int = 5000, seed: int = 0) -> tuple[float, float, float]:
    """Czułość = TP / (TP + FN); TP z puli przepuszczonych (precyzja), FN z puli odrzuconych (odsetek pominiętych).
    Przedział z rozkładów Beta (Jeffreys) obu odsetków."""
    rng = np.random.default_rng(seed)
    p = rng.beta(k_pass + 0.5, n_pass - k_pass + 0.5, draws)
    r = rng.beta(k_rej + 0.5, n_rej - k_rej + 0.5, draws)
    tp, fn = n_pass_pool * p, n_rej_pool * r
    rec = tp / (tp + fn)
    point_tp = n_pass_pool * (k_pass / n_pass)
    point_fn = n_rej_pool * (k_rej / n_rej)
    est = point_tp / (point_tp + point_fn) if (point_tp + point_fn) > 0 else float("nan")
    return est, float(np.quantile(rec, 0.025)), float(np.quantile(rec, 0.975))


def read_reviewed(path: Path) -> pd.DataFrame:
    if path.suffix.lower() in (".xlsx", ".xls"):
        with pd.ExcelFile(path) as xl:                      # zamykamy plik (Windows blokuje otwarty xlsx)
            d = xl.parse("Przeglad" if "Przeglad" in xl.sheet_names else xl.sheet_names[0], dtype=str)
    else:
        d = pd.read_csv(path, sep=None, engine="python", encoding="utf-8-sig", dtype=str)
    d.columns = [str(c).strip().lower() for c in d.columns]
    for c in ("id", "o_uchodzcach"):
        if c not in d.columns:
            raise SystemExit(f"{path.name}: brak kolumny {c} (są: {list(d.columns)})")
    d["label"] = d["o_uchodzcach"].fillna("").str.strip().str.lower().replace({"": None, "tak": "tak", "nie": "nie", "niejasne": "niejasne"})
    bad = d[d["label"].notna() & ~d["label"].isin(CHOICES)]
    if len(bad):
        raise SystemExit(f"{path.name}: nieznane wartości {sorted(bad['label'].unique())}; dozwolone: {CHOICES}")
    return d[["id", "label", "komentarz"] if "komentarz" in d.columns else ["id", "label"]]


def score_md(master: pd.DataFrame, rev: pd.DataFrame, prec_min: float, rec_min: float, out_dir: Path) -> str:
    m = master.merge(rev, left_on="item_id", right_on="id", how="left")
    miss = int(m["label"].isna().sum())
    m = m.dropna(subset=["label"]).copy()
    m["yes"] = m["label"].eq("tak")
    L = ["# Filtr uchodźczy: wynik ręcznego przeglądu", "",
         f"Oceniono {len(m)} z {len(master)} postów (braków: {miss}). „niejasne” liczymy na niekorzyść filtra: w puli przepuszczonych jak „nie”, "
         "w puli odrzuconych jak „tak” (czyli jako możliwy brak).", "",
         "| zbiór | precyzja (przepuszczone = o uchodźcach) | pominięte wśród odrzuconych | szac. czułość | oceniono (przep. / odrz.) |", "|---|---|---|---|---|"]
    tot = {"Np": 0.0, "kp": 0, "np": 0, "Nr": 0.0, "kr": 0, "nr": 0}
    rows = []
    for side in SIDES + ["razem"]:
        d = m if side == "razem" else m[m["side"] == side]
        p, r = d[d["group"] == "przepuszczony"], d[d["group"] == "odrzucony"]
        if p.empty or r.empty:
            continue
        kp, npp = int(p["yes"].sum()), len(p)
        kr = int((r["yes"] | r["label"].eq("niejasne")).sum()); nr = len(r)
        Np = float(p["pool_N"].groupby(p["side"]).first().sum()); Nr = float(r["pool_N"].groupby(r["side"]).first().sum())
        pr, plo, phi = wilson(kp, npp)
        mr, mlo, mhi = wilson(kr, nr)
        rec, rlo, rhi = recall_ci(Np, kp, npp, Nr, kr, nr)
        rows.append((side, pr, plo, rec, rlo, kp, npp, kr, nr))
        L.append(f"| {side} | {pr:.2f} [{plo:.2f}; {phi:.2f}] ({kp}/{npp}) | {mr:.2f} [{mlo:.2f}; {mhi:.2f}] ({kr}/{nr}) | "
                 f"{rec:.2f} [{rlo:.2f}; {rhi:.2f}] | {npp} / {nr} |")
    L += ["", "Przedziały 95% (Wilson dla odsetków; Beta dla czułości, ważona wielkością pul). Przy ~20 postach na komórkę są szerokie. "
          "Czułość opiera się na rzadkich trafieniach w puli odrzuconych, więc jest szacunkiem zgrubnym.", "", "## Wnioski", ""]
    if not rows:
        L.append("- Za mało ocenionych postów, żeby cokolwiek wywnioskować.")
    else:
        side, pr, plo, rec, rlo = rows[-1][0], rows[-1][1], rows[-1][2], rows[-1][3], rows[-1][4]
        if pr < prec_min:
            L.append(f"- **Precyzja {pr:.2f} < {prec_min}:** zbyt wiele przepuszczonych postów nie jest o uchodźcach. Otwórz `filtr_uchodzczy_bledy.csv` "
                     "(typ FP) i usuń lub zawęź rdzenie, które stoją za fałszywymi trafieniami (lista poniżej).")
        else:
            L.append(f"- Precyzja {pr:.2f} ≥ {prec_min}: przepuszczone posty w większości są o uchodźcach.")
        if rec < rec_min:
            L.append(f"- **Szacowana czułość {rec:.2f} < {rec_min}:** filtr gubi dużo postów uchodźczych. Zobacz w `filtr_uchodzczy_bledy.csv` "
                     "teksty typu FN i dopisz brakujące rdzenie (po rosyjsku / ukraińsku, z odmianami).")
        else:
            L.append(f"- Szacowana czułość {rec:.2f} ≥ {rec_min}: filtr nie gubi większości postów uchodźczych (szacunek zgrubny).")
        L.append("- Po każdej zmianie listy podnieś `TELEGRAM_REFUGEE_VERSION` w `config.py`, uruchom `check` ponownie i oceń NOWĄ próbę "
                 "(nie tę samą, na której poprawiałeś listę).")
    err = pd.concat([m[(m["group"] == "przepuszczony") & ~m["yes"]].assign(typ="FP"),
                     m[(m["group"] == "odrzucony") & (m["yes"] | m["label"].eq("niejasne"))].assign(typ="FN")])
    out_dir.mkdir(parents=True, exist_ok=True)
    if len(err):
        err[["item_id", "typ", "side", "label", "refugee_hits", "text_pl", "text_original"]].to_csv(
            out_dir / "filtr_uchodzczy_bledy.csv", index=False, encoding="utf-8-sig")
        fp = Counter(w for h in err.loc[err["typ"] == "FP", "refugee_hits"].fillna("") for w in set(h.split("; ")) if w)
        if fp:
            L += ["", "Fragmenty najczęściej stojące za fałszywymi trafieniami (FP):", "", "| fragment | postów |", "|---|---|"]
            L += [f"| {w} | {c} |" for w, c in fp.most_common(10)]
    return "\n".join(L) + "\n"


# --------------------------------------------------------------------------- CLI
def cmd_check(a) -> None:
    sm.log("Krok 1/4: wczytanie i filtrowanie")
    df = load(a.posts, a.clusters)
    sm.log(f"  przepuszczone przez filtr uchodźczy: {sm.num(df['is_refugee'].sum())} z {sm.num(len(df))}")
    sm.log("Krok 2/4: statystyki gęstości i szereg dzienny")
    idx = pd.date_range(config.EFFECTIVE_START, config.STUDY_END, freq="D")
    stats = {side: side_stats(df[df["side"] == side], idx) for side in SIDES}
    cats = {k: int(df.loc[df["is_refugee"], k].sum()) for k in config.TELEGRAM_REFUGEE}
    top_hits = Counter(w for h in df.loc[df["is_refugee"], "refugee_hits"] for w in set(h.split("; ")) if w).most_common(15)
    md = results_md(stats, cats, top_hits, config.TELEGRAM_REFUGEE_VERSION)
    REPORTS.mkdir(parents=True, exist_ok=True)
    (REPORTS / "filtr_uchodzczy_wyniki.md").write_text(md, encoding="utf-8")
    (REPORTS / "filtr_uchodzczy_summary.json").write_text(json.dumps(stats, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    config.PROCESSED.mkdir(parents=True, exist_ok=True)
    daily_table(df, idx).to_csv(config.PROCESSED / "telegram_refugee_daily.csv", index=False)

    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    sm.log("Krok 3/4: losowanie próby do przeglądu ręcznego")
    master = draw_review(df, a.n_pass, a.n_rej, a.seed)
    sm.log(f"  {len(master)} postów; tłumaczenie ({a.translator})")
    tr = sm.make_translator(a.translator)
    if tr is not None:
        tl, st = sm.translate_all(master["text_original"], tr, out / "translation_cache.json")
        master["text_pl"] = [sm.clip(t) for t in tl]
        master["translation_status"] = st
        bad = sum(s != "ok" for s in st)
        if bad:
            print(f"UWAGA: {bad} tłumaczeń się nie powiodło; uruchom ponownie (udane są w cache).")
    sm.log("Krok 4/4: zapis arkuszy do przeglądu")
    master.to_csv(out / "filtr_przeglad_master.csv", index=False, encoding="utf-8-sig")
    view = review_view(master)
    write_review_xlsx(out / "filtr_przeglad.xlsx", view)
    view.to_csv(out / "filtr_przeglad.csv", index=False, sep=";", encoding="utf-8-sig")
    print(md)
    print(f"Wynik diagnostyki: reports/filtr_uchodzczy_wyniki.md; szereg dzienny: data/processed/telegram_refugee_daily.csv")
    print(f"Do przeglądu ręcznego ({len(master)} postów): {out / 'filtr_przeglad.xlsx'} (kolumna o_uchodzcach: tak / nie / niejasne), potem: "
          f"python -m pipeline.telegram_refugee score")


def cmd_score(a) -> None:
    out = Path(a.out)
    master = pd.read_csv(out / "filtr_przeglad_master.csv", encoding="utf-8-sig", dtype={"item_id": str}) if not a.master else \
        pd.read_csv(a.master, encoding="utf-8-sig", dtype={"item_id": str})
    rev_path = Path(a.reviewed) if a.reviewed else out / "filtr_przeglad.xlsx"
    md = score_md(master, read_reviewed(rev_path), a.precision_min, a.recall_min, out)
    (REPORTS / "filtr_uchodzczy_przeglad.md").write_text(md, encoding="utf-8")
    print(md)
    print("Zapisano reports/filtr_uchodzczy_przeglad.md oraz (jeśli są błędy) annotation_work/filtr_uchodzczy_bledy.csv")


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check", help="diagnostyka gęstości + próba do przeglądu ręcznego")
    c.add_argument("--posts", default=str(config.ROOT / "data" / "telegram" / "posts_2022_part_*.csv"))
    c.add_argument("--clusters", default=str(config.ROOT / "data" / "telegram" / "leiden_clusters.csv"))
    c.add_argument("--n-pass", type=int, default=40, help="ile postów PRZEPUSZCZONYCH do przeglądu (dzielone na strony)")
    c.add_argument("--n-rej", type=int, default=40, help="ile postów ODRZUCONYCH do przeglądu (dzielone na strony)")
    c.add_argument("--translator", choices=["google", "deepl", "argos", "none"], default="google")
    c.add_argument("--seed", type=int, default=42)
    c.add_argument("--out", default=str(ANNOT))
    c.set_defaults(fn=cmd_check)
    s = sub.add_parser("score", help="precyzja i czułość filtra z wypełnionego przeglądu")
    s.add_argument("--master", help="domyślnie annotation_work/filtr_przeglad_master.csv")
    s.add_argument("--reviewed", help="domyślnie annotation_work/filtr_przeglad.xlsx")
    s.add_argument("--precision-min", type=float, default=0.7)
    s.add_argument("--recall-min", type=float, default=0.5)
    s.add_argument("--out", default=str(ANNOT))
    s.set_defaults(fn=cmd_score)
    a = ap.parse_args(argv)
    sm.use_utf8_console()
    a.fn(a)


if __name__ == "__main__":
    main()
